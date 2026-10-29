"""Word invoices: extract tables/paragraphs, replace description text in place, verify.

.doc files are converted to .docx first (LibreOffice if installed, else macOS textutil).

Word lays text out when the file is opened, so a description that wrapped onto five lines
would leave a one-line row once replaced. LibreOffice renders the invoice with the edited
paragraphs (and the rest of their row) in marker colours - colour never changes layout - to
measure them: an edited table row gets its original height as a minimum height, and other
blocks are padded with blank line breaks to their original number of lines. A second render
checks that every other word is still where it was.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips

from .models import ApplyResult, Cell, Layout, Plan, Row, Table, VerificationError

PROBE_RED = 0x0B  # marker colours are (11, 8*(i+1), 8*(j+1))
LAYOUT_TOLERANCE = 1.0  # points a word may move between two renders (rounding noise)
CELL_ID_RE = re.compile(r"(t\d+-r\d+-c)\d+")
VALIGN = {"center": WD_CELL_VERTICAL_ALIGNMENT.CENTER, "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
          "both": WD_CELL_VERTICAL_ALIGNMENT.BOTH}


@dataclass
class Span:
    """One run of rendered text, tagged with the marker key of the paragraph it came from."""
    page: int
    key: int | None
    text: str
    x: float
    x1: float
    y: float  # baseline
    size: float


@dataclass
class Render:
    spans: list[Span]
    pages: int
    plain_pdf: Path  # as is, for previews
    probe_pdf: Path  # with marker colours


def extract_layout(path: Path, with_images: bool = False) -> Layout:
    layout, _ = _extract(Document(str(path)))
    return layout


def apply_plan(src: Path, plan: Plan, dst: Path, text: str, preview_dir: Path | None) -> ApplyResult:
    if not soffice():
        raise VerificationError("LibreOffice is needed to check the layout of Word invoices")
    doc = Document(str(src))
    _, elements = _extract(doc)
    before = _snapshot(doc, elements)

    groups: list[list[str]] = []
    for group in plan.edits:
        unknown = [b for b in group.block_ids if b not in elements]
        if unknown:
            raise VerificationError(f"plan refers to unknown blocks {unknown}")
        if any(_text(elements[b]).strip() for b in group.block_ids):
            groups.append(group.block_ids)
    probes = _probe_keys(elements, groups)

    result = ApplyResult(fonts=["(unchanged: Word keeps each run's own font)"])
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp = Path(tmp_dir)
        rb = _render(src, probes, tmp / "before")
        cells = [_keys(probes, ids[:1]) for ids in groups]
        heights = _row_heights(rb, cells)
        for ids, height in zip(groups, heights):
            _replace(elements, ids, probes, rb.spans, text, height)
        edited = tmp / "edited.docx"
        doc.save(str(edited))
        ra = _render(edited, probes, tmp / "after")

        # LibreOffice draws a row a little taller than its height setting (by about its
        # borders), so measure once more and take off what was added.
        drift = [(ids, h, a) for ids, h, a in zip(groups, heights, _row_heights(ra, cells))
                 if h and a and abs(a - h) > 0.2 and elements[ids[0]].tag == qn("w:tc")]
        if drift:
            for ids, h, a in drift:
                _keep_row_height(elements[ids[0]].getparent(), h - (a - h), force=True)
            doc.save(str(edited))
            ra = _render(edited, probes, tmp / "after-corrected")

        after_doc = Document(str(edited))
        _, after_elements = _extract(after_doc)
        _verify(before, _snapshot(after_doc, after_elements), groups, text)
        _verify_layout(rb, ra, [_keys(probes, ids) for ids in groups], text)

        shutil.copyfile(edited, dst)
        if preview_dir:
            from .pdf_doc import render_compare
            result.previews = render_compare(rb.plain_pdf, ra.plain_pdf, preview_dir / src.stem)
    return result


# --- extraction ---------------------------------------------------------------------------

def _extract(doc) -> tuple[Layout, dict]:
    tables: list[Table] = []
    elements: dict = {}
    for ti, table in enumerate(doc.tables):
        rows = []
        width = 0
        for ri, tr in enumerate(table._tbl.tr_lst):
            col = tr.grid_before or 0
            cells = []
            for ci, tc in enumerate(tr.tc_lst):
                cid = f"t{ti}-r{ri}-c{ci}"
                elements[cid] = tc
                text = "" if tc.vMerge == "continue" else _text(tc)
                cells.append(Cell(cid, text, col, col + tc.grid_span))
                col += tc.grid_span
            width = max(width, col)
            rows.append(Row(cells))
        tables.append(Table(f"t{ti}", 0, 0, width, rows, page_flow=False))

    free = []
    for pi, p in enumerate(doc.paragraphs):
        pid = f"p{pi}"
        elements[pid] = p._p
        if _text(p._p).strip():
            free.append(Cell(pid, _text(p._p)))
    full = "\n".join([c.text for t in tables for r in t.rows for c in r.cells] + [c.text for c in free])
    return Layout(tables, free, full_text=full), elements


def _text(el) -> str:
    return "\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in el.iter(qn("w:p")))


# --- replacement --------------------------------------------------------------------------

def _replace(elements: dict, ids: list[str], probes: dict, spans: list[Span], text: str,
             row_height: float | None) -> None:
    """Put the text in the longest run of the first line and blank the rest of the block.

    A table row keeps its rendered height as a minimum height, so each edited cell keeps just
    the paragraph holding the text (paragraphs with pictures stay). When the row's amount sits
    lower than the cell's first line, a top-aligned cell takes the amount cell's vertical
    alignment (centre/bottom), or else moves down with leading line breaks. Other blocks are
    padded with line breaks so every paragraph keeps the number of lines it had.
    """
    paragraphs = [(probes[(b, i)], p) for b in ids for i, p in enumerate(elements[b].iter(qn("w:p")))]
    texted = [(key, p) for key, p in paragraphs if _text(p).strip()]
    lines = {key: _lines(spans, key) for key, _ in texted}

    lead_key, lead = texted[0]
    keep = max(lead.iter(qn("w:t")), key=lambda t: len((t.text or "").strip()))
    for _, p in paragraphs:
        for t in p.iter(qn("w:t")):
            t.text = ""  # runs and paragraphs are kept
        numbering = p.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
        if numbering is not None:  # a bullet/number would still show in front of the text
            numbering.getparent().remove(numbering)
    keep.text = text

    lead_lines = lines[lead_key]
    cell = elements[ids[0]]
    anchor = _row_anchor(ids, probes, spans)
    above = _anchor_line(lead_lines, anchor[1:] if anchor else None)
    if row_height and all(elements[b].tag == qn("w:tc") for b in ids):
        _keep_row_height(cell.getparent(), row_height)
        for b in ids:
            cell_paragraphs = [p for p in elements[b] if p.tag == qn("w:p")]
            stay = lead if lead in cell_paragraphs else cell_paragraphs[0]  # a cell needs one
            for p in cell_paragraphs:
                if p is not stay and not any(p.iter(qn("w:drawing"), qn("w:pict"), qn("w:object"))):
                    p.getparent().remove(p)
        if above and _valign(cell) == "top":
            if (peer := _valign(elements[anchor[0]])) in VALIGN:
                cell.get_or_add_tcPr().vAlign_val = VALIGN[peer]
            else:
                for _ in range(above):
                    keep.addprevious(OxmlElement("w:br"))
        return

    for _ in range(above):
        keep.addprevious(OxmlElement("w:br"))
    for _ in range(len(lead_lines) - 1 - above):
        keep.addnext(OxmlElement("w:br"))
    for key, p in texted[1:]:
        first = next(p.iter(qn("w:t")))
        for _ in range(len(lines[key]) - 1):
            first.addnext(OxmlElement("w:br"))


def _valign(tc) -> str:
    v = tc.tcPr.find(qn("w:vAlign")) if tc.tcPr is not None else None
    return "top" if v is None else v.get(qn("w:val"), "top")


def _keep_row_height(tr, height_pt: float, force: bool = False) -> None:
    """Give the row its rendered height as a minimum ("at least") height."""
    current = tr.trPr.find(qn("w:trHeight")) if tr.trPr is not None else None
    if current is not None and current.get(qn("w:hRule")) == "exact":
        return
    twips = round(height_pt * 20)
    if (not force and current is not None and current.get(qn("w:hRule")) == "atLeast"
            and int(current.get(qn("w:val"), 0)) >= twips):
        return
    tr.trHeight_val = Twips(twips)
    tr.trHeight_hRule = WD_ROW_HEIGHT_RULE.AT_LEAST


def _row_heights(render: Render, cells: list[set[int]]) -> list[float | None]:
    """Rendered height of the table row holding each cell (None when it can't be measured)."""
    heights: list[float | None] = []
    tables: dict[int, list] = {}
    with pymupdf.open(render.probe_pdf) as pdf:
        for keys in cells:
            first = min((s for s in render.spans if s.key in keys), key=lambda s: (s.page, s.y), default=None)
            height = None
            if first:
                if first.page not in tables:
                    tables[first.page] = pdf[first.page].find_tables(use_layout=False).tables
                point = pymupdf.Point(first.x + 0.5, first.y - 1)
                for tab in tables[first.page]:
                    for row in tab.rows:
                        if any(c and pymupdf.Rect(c).contains(point) for c in row.cells):
                            height = row.bbox[3] - row.bbox[1]
            heights.append(height)
    return heights


def _lines(spans: list[Span], key: int) -> list[tuple[int, float]]:
    """(page, baseline) of each rendered line of a paragraph; superscripts join their line."""
    lines: list[list[Span]] = []
    for s in sorted((s for s in spans if s.key == key), key=lambda s: (s.page, s.y)):
        last = lines[-1] if lines else None
        if last and last[0].page == s.page and s.y - last[0].y <= 0.6 * max(t.size for t in last + [s]):
            last.append(s)
        else:
            lines.append([s])
    return [(line[0].page, max(line, key=lambda s: s.size).y) for line in lines] or [(0, 0.0)]


def _row_anchor(ids: list[str], probes: dict, spans: list[Span]) -> tuple[str, int, float] | None:
    """(block id, page, first baseline) of the rightmost other cell in the row (normally the amount)."""
    m = CELL_ID_RE.fullmatch(ids[0])
    if not m:
        return None
    peers: dict[str, list[Span]] = {}
    for (block, _), key in probes.items():
        if block.startswith(m.group(1)) and block not in ids:
            peers.setdefault(block, []).extend(s for s in spans if s.key == key)
    peers = {b: ss for b, ss in peers.items() if ss}
    if not peers:
        return None
    block = max(peers, key=lambda b: min(s.x for s in peers[b]))
    first = min(peers[block], key=lambda s: (s.page, s.y))
    return block, first.page, first.y


def _anchor_line(lines: list[tuple[int, float]], anchor: tuple[int, float] | None) -> int:
    """Index of the line to write on: the top line, unless the row's amount sits lower
    inside this block (vertically centred rows), as in the PDF path."""
    if not anchor or len(lines) < 2:
        return 0
    page, y = anchor
    first, last = lines[0], lines[-1]
    if first[0] == page and first[1] + 2 < y <= last[1] + 0.5:
        return min(range(len(lines)), key=lambda i: abs(lines[i][1] - y) if lines[i][0] == page else 1e9)
    return 0


# --- rendering ----------------------------------------------------------------------------

def _probe_keys(elements: dict, groups: list[list[str]]) -> dict[tuple[str, int], int]:
    """Marker key for every paragraph of the edited blocks and of the other cells in their rows."""
    blocks: list[str] = []
    for ids in groups:
        blocks += ids
        if m := CELL_ID_RE.fullmatch(ids[0]):
            blocks += [b for b in elements if b.startswith(m.group(1)) and b not in ids]
    probes: dict[tuple[str, int], int] = {}
    for b in dict.fromkeys(blocks):
        for i, _ in enumerate(elements[b].iter(qn("w:p"))):
            probes[(b, i)] = len(probes)
    if len(probes) > 31 * 31:
        raise VerificationError("too many paragraphs to measure")
    return probes


def _keys(probes: dict, ids: list[str]) -> set[int]:
    return {k for (b, _), k in probes.items() if b in ids}


def _probe_color(key: int) -> str:
    g, b = divmod(key, 31)
    return f"{PROBE_RED:02X}{8 * (g + 1):02X}{8 * (b + 1):02X}"


def _probe_key(rgb: int) -> int | None:
    r, g, b = rgb >> 16 & 255, rgb >> 8 & 255, rgb & 255
    gi, bi = round(g / 8) - 1, round(b / 8) - 1
    if abs(r - PROBE_RED) > 1 or gi < 0 or bi < 0 or abs(g - 8 * (gi + 1)) > 1 or abs(b - 8 * (bi + 1)) > 1:
        return None
    return gi * 31 + bi


def _render(src: Path, probes: dict, out_dir: Path) -> Render:
    """Render src with LibreOffice as is (for previews) and with marker colours (to measure)."""
    out_dir.mkdir(parents=True)
    plain = out_dir / "plain.docx"
    shutil.copyfile(src, plain)
    doc = Document(str(src))
    _, elements = _extract(doc)
    paragraphs = {b: list(elements[b].iter(qn("w:p"))) for b, _ in probes}
    for (b, i), key in probes.items():
        if i >= len(paragraphs[b]):
            continue  # a blank paragraph removed from an edited cell
        para = paragraphs[b][i]
        for run in para.iter(qn("w:r")):
            color = run.get_or_add_rPr().get_or_add_color()
            for attr in ("w:themeColor", "w:themeTint", "w:themeShade"):
                color.attrib.pop(qn(attr), None)
            color.set(qn("w:val"), _probe_color(key))
        # the paragraph mark's colour also colours its bullet/number label
        ppr = para.get_or_add_pPr()
        mark = ppr.find(qn("w:rPr"))
        if mark is None:
            mark = OxmlElement("w:rPr")
            tail = next((c for c in ppr if c.tag in (qn("w:sectPr"), qn("w:pPrChange"))), None)
            tail.addprevious(mark) if tail is not None else ppr.append(mark)
        for old in mark.findall(qn("w:color")):
            mark.remove(old)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), _probe_color(key))
        mark.append(color)
    doc.save(str(out_dir / "probe.docx"))
    subprocess.run([soffice(), "--headless", "--convert-to", "pdf", "--outdir", str(out_dir),
                    str(plain), str(out_dir / "probe.docx")], check=True, capture_output=True, timeout=180)
    with pymupdf.open(out_dir / "probe.pdf") as pdf:
        spans = [
            Span(page.number, _probe_key(s["color"]), s["text"].strip(), s["origin"][0], s["bbox"][2],
                 s["origin"][1], s["size"])
            for page in pdf
            for block in page.get_text("dict")["blocks"]
            for line in block.get("lines", [])
            for s in line["spans"] if s["text"].strip()
        ]
        return Render(spans, len(pdf), out_dir / "plain.pdf", out_dir / "probe.pdf")


# --- verification -------------------------------------------------------------------------

def _snapshot(doc, elements: dict) -> dict:
    body = doc.element.body
    # Read header/footer parts directly: section.header/.footer would create missing ones.
    parts = [r.target_part for r in doc.part.rels.values() if r.reltype in (RT.HEADER, RT.FOOTER)]
    return {
        "texts": {k: _text(el) for k, el in elements.items()},
        "shape": [len(t.rows) for t in doc.tables] + [len(doc.paragraphs)],
        "drawings": len(list(body.iter(qn("w:drawing")))),
        "headers": sorted(_text(p.element) for p in parts),
    }


def _verify(before: dict, after: dict, edited: list[list[str]], text: str) -> None:
    for key in ("shape", "drawings", "headers"):
        if before[key] != after[key]:
            raise VerificationError(f"document {key} changed")
    touched = {b for group in edited for b in group}
    for k, v in before["texts"].items():
        if k not in touched and after["texts"].get(k) != v:
            raise VerificationError(f"text outside the edits changed in {k}: {v!r} -> {after['texts'].get(k)!r}")
    for group in edited:
        got = " ".join(after["texts"][b].strip() for b in group if after["texts"][b].strip())
        if got != text:
            raise VerificationError(f"edited block {group[0]} reads {got!r}, expected {text!r}")


def _verify_layout(before: Render, after: Render, groups: list[set[int]], text: str) -> None:
    if before.pages != after.pages:
        raise VerificationError(f"page count changed {before.pages} -> {after.pages}")
    edited = set().union(*groups)
    outside_before = [s for s in before.spans if s.key not in edited]
    outside_after = [s for s in after.spans if s.key not in edited]
    lost, added = _unmatched(outside_before, outside_after), _unmatched(outside_after, outside_before)
    if lost or added:
        def show(spans):
            return [(s.text, s.page + 1, round(s.x), round(s.y)) for s in spans[:4]]
        raise VerificationError(f"text outside the edits moved; before={show(lost)} after={show(added)}")
    for keys in groups:
        spans = sorted((s for s in after.spans if s.key in keys), key=lambda s: (s.page, s.y, s.x))
        got = " ".join(s.text for s in spans)
        if " ".join(got.split()) != text:
            raise VerificationError(f"edited block renders as {got!r}, expected {text!r}")
        for a, b in zip(spans, spans[1:]):  # justified text stretched across the cell
            if b.x - a.x1 > a.size:
                raise VerificationError(f"edited text is stretched ({b.x - a.x1:.0f}pt gap)")


def _unmatched(spans: list[Span], others: list[Span]) -> list[Span]:
    """Spans with no counterpart (same page and text, within LAYOUT_TOLERANCE) in others."""
    pool: dict[tuple[int, str], list[Span]] = {}
    for s in others:
        pool.setdefault((s.page, s.text), []).append(s)
    missing = []
    for s in spans:
        candidates = pool.get((s.page, s.text), [])
        hit = next((c for c in candidates
                    if abs(c.x - s.x) <= LAYOUT_TOLERANCE and abs(c.y - s.y) <= LAYOUT_TOLERANCE), None)
        if hit:
            candidates.remove(hit)
        else:
            missing.append(s)
    return missing


# --- .doc conversion ----------------------------------------------------------------------

def verify_roundtrip(original: Path, final: Path, removed: list[str], groups: int, text: str,
                     work_dir: Path, preview_prefix: Path | None) -> None:
    """End-to-end check for .doc files, whose trip through .docx can move things: rendered
    side by side, every word must still be in place except the replaced description words."""
    before = convert(original, "pdf", work_dir / "original")
    after = convert(final, "pdf", work_dir / "final")
    with pymupdf.open(before) as a, pymupdf.open(after) as b:
        if len(a) != len(b):
            raise VerificationError(f"page count changed {len(a)} -> {len(b)}")
        wa = Counter((w[4], p.number, round(w[0]), round(w[1])) for p in a for w in p.get_text("words"))
        wb = Counter((w[4], p.number, round(w[0]), round(w[1])) for p in b for w in p.get_text("words"))
    lost = Counter(w[0] for w in (wa - wb).elements()) - Counter(" ".join(removed).split())
    added = Counter(w[0] for w in (wb - wa).elements()) - Counter(text.split() * groups)
    if lost or added:
        raise VerificationError(f"words moved or changed; lost={list(lost)[:5]} added={list(added)[:5]}")
    if preview_prefix:
        from .pdf_doc import render_compare
        render_compare(before, after, preview_prefix)

def soffice() -> str | None:
    for candidate in (shutil.which("soffice"), shutil.which("libreoffice"),
                      "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if candidate and Path(candidate).exists():
            return candidate
    return None


def convert(src: Path, fmt: str, out_dir: Path) -> Path:
    """Convert with LibreOffice (any direction) or macOS textutil (doc <-> docx only)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{src.stem}.{fmt}"
    if office := soffice():
        subprocess.run([office, "--headless", "--convert-to", fmt, "--outdir", str(out_dir), str(src)],
                       check=True, capture_output=True, timeout=180)
    elif shutil.which("textutil") and fmt in ("docx", "doc"):
        subprocess.run(["textutil", "-convert", fmt, "-output", str(out), str(src)],
                       check=True, capture_output=True, timeout=180)
    else:
        raise RuntimeError(f"cannot convert {src.name} to .{fmt}: install LibreOffice")
    if not out.exists():
        raise RuntimeError(f"conversion of {src.name} to .{fmt} produced no file")
    return out
