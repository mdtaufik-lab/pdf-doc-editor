"""Read the Relook list, decide exclusions, and match invoice files to their sheet row."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .config import (
    COMPANY_SUFFIX_RE, EXCLUDED_COMPANIES, EXCLUDED_TYPES, RELOOK_SHEET_NAME,
)

DRIVE_ID_RE = re.compile(r"(?:id=|/d/)([\w-]{20,})")
TITLE_WORDS = {"mr", "mrs", "ms", "miss", "dr", "adv", "advocate", "prof", "shri", "smt"}


@dataclass
class RelookRow:
    sheet_row: int  # 1-based row number in Excel
    fy: str
    month: str
    name: str
    category: str
    description: str
    type: str
    invoice_no: str
    invoice_date: str
    invoice_amount: str
    amount_paid: str
    link: str

    @property
    def drive_id(self) -> str | None:
        m = DRIVE_ID_RE.search(self.link)
        return m.group(1) if m else None


def load_relook(path: Path) -> list[RelookRow]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = next((ws for ws in wb.worksheets if ws.title.strip().lower() == RELOOK_SHEET_NAME), None)
    if ws is None:
        raise ValueError(f"No '{RELOOK_SHEET_NAME}' sheet in {path.name}: {wb.sheetnames}")
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h or "").strip().lower() for h in rows[0]]

    def col(*names: str) -> int:
        for i, h in enumerate(header):
            if any(h.startswith(n) for n in names):
                return i
        raise ValueError(f"Relook list is missing a column starting with {names}")

    idx = {
        "fy": col("fy"), "month": col("month"), "name": col("name of person"),
        "category": col("category"), "description": col("description"), "type": col("type"),
        "invoice_no": col("invoice no"), "invoice_date": col("invoice date"),
        "invoice_amount": col("invoice amount"), "amount_paid": col("amount paid"),
        "link": col("invoice link"),
    }
    out = []
    for n, r in enumerate(rows[1:], start=2):
        if not any(r):
            continue
        out.append(RelookRow(sheet_row=n, **{k: _s(r[i]) for k, i in idx.items()}))
    return out


def exclusion_reason(row: RelookRow) -> str | None:
    if row.type.strip().lower() in EXCLUDED_TYPES:
        return f"type: {row.type}"
    name = row.name.lower()
    if any(c in name for c in EXCLUDED_COMPANIES) or COMPANY_SUFFIX_RE.search(row.name):
        return f"company vendor: {row.name}"
    return None


def match_row(rows: list[RelookRow], file: Path, doc_text: str) -> tuple[RelookRow | None, str]:
    """Find the Relook row for a file. Returns (row, explanation); row is None when unsure.

    Downloaded invoices live in Invoices/<drive id>/<original name>, so the id is also
    looked for in the parent folder's name.
    """
    for row in rows:
        if row.drive_id and row.drive_id in file.name:
            return row, f"drive id {row.drive_id} in filename"
        if row.drive_id and row.drive_id == file.parent.name:
            return row, f"drive id {row.drive_id} in folder name"

    fname = _norm(file.stem)
    text = _norm(doc_text)
    numbers = {n.replace(",", "").lstrip("0") for n in re.findall(r"\d[\d,]*", doc_text)}
    scored = []
    for row in rows:
        score, why = _score(row, fname, text, numbers)
        if score > 0:
            scored.append((score, why, row))
    if not scored:
        return None, "no row with this person's name"
    scored.sort(key=lambda s: s[0], reverse=True)
    best = scored[0]
    if len(scored) > 1 and scored[1][0] == best[0]:
        return None, f"ambiguous: sheet rows {best[2].sheet_row} and {scored[1][2].sheet_row} score equally"
    if best[0] < 4:
        return None, f"weak match (score {best[0]}) to row {best[2].sheet_row}"
    return best[2], f"score {best[0]}: {', '.join(best[1])}"


def _score(row: RelookRow, fname: str, text: str, numbers: set[str]) -> tuple[float, list[str]]:
    full = _norm(row.name)
    tokens = [t for t in full.split() if len(t) >= 3 and t not in TITLE_WORDS]
    if not tokens:
        return 0, []
    score, why = 0.0, []
    if f" {full} " in f" {fname} ":
        score, why = score + 5, why + ["full name in filename"]
    elif f" {full} " in f" {text} ":
        score, why = score + 4, why + ["full name in document"]
    hits = [t for t in tokens if f" {t} " in f" {fname} " or f" {t} " in f" {text} "]
    if not hits:
        return 0, []
    score += len(hits)
    inv = _norm(row.invoice_no)
    if inv and inv not in {"not mentioned", "na"} and re.search(rf"(?<!\w){re.escape(inv)}(?!\w)", f"{fname} {text}"):
        score, why = score + 2, why + [f"invoice no {row.invoice_no}"]
    amounts = {re.sub(r"\D", "", a.split(".")[0]).lstrip("0") for a in (row.amount_paid, row.invoice_amount)}
    if amounts & numbers - {""}:
        score, why = score + 2, why + ["amount"]
    return score, why


def _norm(s: str) -> str:
    return " ".join(re.sub(r"[^a-z0-9/]+", " ", s.lower()).split())


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()
